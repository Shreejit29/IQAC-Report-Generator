from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Iterable
import zipfile

import fitz  # PyMuPDF
from docx import Document
from PIL import Image
from openpyxl import load_workbook


@dataclass
class ExtractedDocument:
    name: str
    text: str
    kind: str
    needs_vision: bool = False


def extract_docx_text_and_images(data: bytes) -> tuple[str, list[Image.Image]]:
    doc = Document(BytesIO(data))
    chunks: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            chunks.append(t)
    for table in doc.tables:
        for row in table.rows:
            vals = [cell.text.strip().replace("\n", " | ") for cell in row.cells]
            if any(vals):
                chunks.append(" | ".join(vals))

    images: list[Image.Image] = []
    try:
        with zipfile.ZipFile(BytesIO(data)) as zf:
            media_names = [name for name in zf.namelist() if name.startswith("word/media/")]
            for name in media_names[:12]:
                raw = zf.read(name)
                try:
                    images.append(Image.open(BytesIO(raw)).convert("RGB"))
                except Exception:
                    continue
    except Exception:
        pass
    return "\n".join(chunks), images


def extract_xlsx_text(data: bytes) -> str:
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    chunks: list[str] = []
    for ws in wb.worksheets:
        chunks.append(f"[Workbook Sheet: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v).strip() if v is not None else "" for v in row]
            if any(vals):
                chunks.append(" | ".join(vals))
    return "\n".join(chunks)


def extract_pdf_text(data: bytes) -> tuple[str, bool, list[Image.Image]]:
    doc = fitz.open(stream=data, filetype="pdf")
    text_parts: list[str] = []
    sparse_pages: list[Image.Image] = []
    for page_idx, page in enumerate(doc):
        txt = page.get_text("text").strip()
        if txt:
            text_parts.append(f"[PDF Page {page_idx + 1}]\n{txt}")
        elif page_idx < 8:
            pix = page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4), alpha=False)
            sparse_pages.append(Image.open(BytesIO(pix.tobytes("png"))).convert("RGB"))
    doc.close()
    text = "\n\n".join(text_parts)
    return text, len(text) < 120, sparse_pages


def extract_document(name: str, data: bytes) -> tuple[ExtractedDocument, list[Image.Image]]:
    suffix = Path(name).suffix.lower()
    if suffix == ".docx":
        text, images = extract_docx_text_and_images(data)
        return ExtractedDocument(name, text, "DOCX", len(text) < 120), images
    if suffix == ".pdf":
        text, sparse, images = extract_pdf_text(data)
        return ExtractedDocument(name, text, "PDF", sparse), images
    if suffix in {".txt", ".md"}:
        return ExtractedDocument(name, data.decode("utf-8", errors="replace"), "TEXT", False), []
    if suffix in {".xlsx", ".xlsm"}:
        text = extract_xlsx_text(data)
        return ExtractedDocument(name, text, "XLSX", len(text) < 120), []
    if suffix in {".jpg", ".jpeg", ".png", ".webp"}:
        return ExtractedDocument(name, "", "IMAGE", True), [Image.open(BytesIO(data)).convert("RGB")]
    raise ValueError(f"Unsupported document type: {suffix}")
