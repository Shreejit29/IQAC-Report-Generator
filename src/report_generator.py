from __future__ import annotations

import io
import os
import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, ImageOps, ImageDraw, ImageFont


FONT_NAME = "Times New Roman"
FONT_SIZE = 12


def _clean(value: object) -> str:
    s = str(value or "").strip()
    return s if s else "Not Identified"


def _set_run_font(run, bold: bool | None = None) -> None:
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(FONT_SIZE)
    if bold is not None:
        run.bold = bold


def _clear_runs(paragraph) -> None:
    for run in paragraph.runs:
        run.text = ""


def _replace_paragraph_value(paragraph, label: str, value: str) -> None:
    _clear_runs(paragraph)
    text = f"{label} {_clean(value)}"
    if paragraph.runs:
        paragraph.runs[0].text = text
        _set_run_font(paragraph.runs[0])
    else:
        run = paragraph.add_run(text)
        _set_run_font(run)


def _replace_table_cell_after_label(cell, value: str) -> None:
    p = cell.paragraphs[0]
    text = p.text
    if text and ":" in text:
        label = text.split(":", 1)[0] + ":"
        _clear_runs(p)
        text_out = f"{label} {_clean(value)}"
        if p.runs:
            p.runs[0].text = text_out
            _set_run_font(p.runs[0])
        else:
            run = p.add_run(text_out)
            _set_run_font(run)
    else:
        _set_cell_text(cell, value)


def _set_cell_text(cell, text: str) -> None:
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    _clear_runs(p)
    value = _clean(text)
    if p.runs:
        p.runs[0].text = value
        _set_run_font(p.runs[0])
    else:
        run = p.add_run(value)
        _set_run_font(run)


def _set_cell_raw(cell, text: str) -> None:
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    _clear_runs(p)
    if p.runs:
        p.runs[0].text = text
        _set_run_font(p.runs[0])
    elif text:
        run = p.add_run(text)
        _set_run_font(run)


def _clear_shape_text(doc: Document, shape_name: str) -> None:
    for element in doc._element.iter():
        if element.tag.endswith("docPr") and element.get("name") == shape_name:
            anchor = element.getparent()
            for child in list(anchor.iter()):
                if child.tag.endswith("t"):
                    child.text = ""


def _insert_picture_fill(doc: Document, shape_name: str, image_path: str) -> None:
    rid, _ = doc.part.get_or_add_image(image_path)
    for element in doc._element.iter():
        if element.tag.endswith("docPr") and element.get("name") == shape_name:
            anchor = element.getparent()
            wsp_nodes = anchor.xpath('.//*[local-name()="wsp"]')
            if not wsp_nodes:
                continue
            wsp = wsp_nodes[0]
            sppr_nodes = wsp.xpath('./*[local-name()="spPr"]')
            if not sppr_nodes:
                continue
            sppr = sppr_nodes[0]
            for child in list(sppr):
                if child.tag.split("}")[-1] in {"noFill", "solidFill", "gradFill", "pattFill", "blipFill", "grpFill"}:
                    sppr.remove(child)
            blip_fill = OxmlElement("a:blipFill")
            blip = OxmlElement("a:blip")
            blip.set(qn("r:embed"), rid)
            blip_fill.append(blip)
            stretch = OxmlElement("a:stretch")
            stretch.append(OxmlElement("a:fillRect"))
            blip_fill.append(stretch)
            lines = sppr.xpath('./*[local-name()="ln"]')
            if lines:
                sppr.insert(sppr.index(lines[0]), blip_fill)
            else:
                sppr.append(blip_fill)
            _clear_shape_text(doc, shape_name)
            return
    raise ValueError(f"Template shape not found: {shape_name}")


def _make_montage(image_paths: list[str], out_path: str, width: int, height: int, title: str) -> str:
    canvas = Image.new("RGB", (width, height), "white")
    count = min(len(image_paths), 6)
    if count == 0:
        raise ValueError("No images supplied")
    if count == 1:
        cols, rows = 1, 1
    elif count == 2:
        cols, rows = 2, 1
    elif count <= 4:
        cols, rows = 2, 2
    else:
        cols, rows = 3, 2
    title_band = max(70, int(height * 0.16))
    pad = max(8, width // 100)
    body_h = height - title_band
    tile_w = max(1, (width - pad * (cols + 1)) // cols)
    tile_h = max(1, (body_h - pad * (rows + 1)) // rows)
    for idx, path in enumerate(image_paths[:count]):
        with Image.open(path) as src:
            src = ImageOps.exif_transpose(src).convert("RGB")
            tile = ImageOps.fit(src, (tile_w, tile_h), method=Image.Resampling.LANCZOS, centering=(0.5, 0.5))
        x = pad + (idx % cols) * (tile_w + pad)
        y = title_band + pad + (idx // cols) * (tile_h + pad)
        canvas.paste(tile, (x, y))
    draw = ImageDraw.Draw(canvas)
    draw.line((0, title_band - 2, width, title_band - 2), fill=(40, 40, 40), width=2)
    font = None
    for candidate in [
        "/usr/share/fonts/truetype/msttcorefonts/Times_New_Roman_Bold.ttf",
        "/usr/share/fonts/truetype/tinos/Tinos-Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf",
    ]:
        if os.path.exists(candidate):
            font = ImageFont.truetype(candidate, max(22, int(width / 52)))
            break
    if font is None:
        font = ImageFont.load_default()
    box = draw.textbbox((0, 0), title, font=font)
    tx = (width - (box[2] - box[0])) / 2
    ty = (title_band - (box[3] - box[1])) / 2 - box[1]
    draw.text((tx, ty), title, fill="black", font=font)
    canvas.save(out_path, format="PNG", optimize=True)
    return out_path


def _prepare_images(image_bytes: Iterable[bytes], work_dir: Path) -> list[str]:
    work_dir.mkdir(parents=True, exist_ok=True)
    paths: list[str] = []
    for i, data in enumerate(image_bytes, 1):
        try:
            img = Image.open(io.BytesIO(data))
            img = ImageOps.exif_transpose(img).convert("RGB")
            p = work_dir / f"photo_{i}.jpg"
            img.save(p, "JPEG", quality=94)
            paths.append(str(p))
        except Exception:
            continue
    return paths


def _safe_filename(title: str, suffix: str) -> str:
    base = re.sub(r"[^A-Za-z0-9 _-]+", "", title or "IQAC_Activity").strip() or "IQAC_Activity"
    base = re.sub(r"\s+", "_", base)
    return f"{base}_{suffix}.docx"


def _set_run_text_with_font(paragraph, text: str) -> None:
    _clear_runs(paragraph)
    run = paragraph.add_run(text)
    _set_run_font(run)


def generate_event_report(
    template_path: str,
    data: dict,
    geotagged: list[bytes],
    normal: list[bytes],
    news: list[bytes],
    work_dir: Path,
) -> tuple[bytes, str]:
    doc = Document(template_path)
    title = _clean(data.get("title"))
    dt_parts = [x for x in [_clean(data.get("activity_date")), _clean(data.get("time"))] if x != "Not Identified"]
    dt = " / ".join(dt_parts) if dt_parts else "Not Identified"
    venue = _clean(data.get("venue"))

    for idx, value in [(0, title), (1, dt), (2, venue), (13, title), (14, dt), (15, venue)]:
        p = doc.paragraphs[idx]
        label = p.text.strip().split(":", 1)[0] + ":"
        _replace_paragraph_value(p, label, value)

    report_text = _clean(data.get("event_report"))
    cell = doc.tables[0].cell(0, 0)
    _set_run_text_with_font(cell.paragraphs[0], report_text)

    work_dir.mkdir(parents=True, exist_ok=True)
    if geotagged:
        paths = _prepare_images(geotagged, work_dir / "geo")
        if paths:
            geo_montage = _make_montage(paths, str(work_dir / "geotagged.png"), 1600, 500, "Geotagged Photo/s")
            _insert_picture_fill(doc, "Rectangle 1", geo_montage)
    if normal:
        paths = _prepare_images(normal, work_dir / "normal")
        if paths:
            normal_montage = _make_montage(paths, str(work_dir / "normal.png"), 1600, 500, "Normal Photos")
            _insert_picture_fill(doc, "Rectangle 15", normal_montage)
    if news:
        paths = _prepare_images(news, work_dir / "news")
        if paths:
            news_montage = _make_montage(paths, str(work_dir / "news.png"), 1500, 620, "News Photo")
            _insert_picture_fill(doc, "Rectangle 13", news_montage)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), _safe_filename(title, "Event_Report")


def generate_iqac_summary(template_path: str, data: dict, proof_flags: dict[str, bool]) -> tuple[bytes, str]:
    doc = Document(template_path)
    top = doc.tables[0]
    _replace_table_cell_after_label(top.cell(0, 0), data.get("iqac_cell_activity_number"))
    _replace_table_cell_after_label(top.cell(1, 0), data.get("activity_schedule_number"))
    _replace_table_cell_after_label(top.cell(2, 0), data.get("activity_type"))

    _replace_paragraph_value(doc.paragraphs[4], "Title of the Activity:", data.get("title"))
    _replace_paragraph_value(doc.paragraphs[5], "Brief Information about the Activity:", data.get("brief_information"))

    fields = [
        "proposal_date", "activity_date", "time", "faculty", "department_committee_association",
        "venue", "participants", "activity_for_class_group", "coordinator_name_phone", "members_support",
        "invited_guest",
    ]
    table1 = doc.tables[1]
    for row, key in enumerate(fields):
        value = data.get(key)
        if key == "invited_guest" and _clean(data.get("resource_person")) != "Not Identified":
            value = data.get("resource_person")
        _set_cell_text(table1.cell(row, 1), value)

    fields2 = ["profile_topic", "objective", "methodology", "outcome"]
    table2 = doc.tables[2]
    for row, key in enumerate(fields2):
        _set_cell_text(table2.cell(row, 1), data.get(key))
    table2.rows[3].height = Inches(0.60)
    table2.rows[3].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    table3 = doc.tables[3]
    mapping = [
        (1, "proposal"),
        (2, "notice"),
        (3, "programme"),
        (4, "invitation"),
        (5, "attendance"),
        (6, "event_report_evidence"),
        (7, "feedback"),
        (8, "news"),
        (9, "publicity"),
        (10, "certificate"),
        (11, "other"),
    ]
    for row, key in mapping:
        _set_cell_raw(table3.cell(row, 2), "✓" if proof_flags.get(key, False) else "")

    out = io.BytesIO()
    doc.save(out)
    title = _safe_filename(_clean(data.get("title")), "IQAC_Summary_Report")
    return out.getvalue(), title
